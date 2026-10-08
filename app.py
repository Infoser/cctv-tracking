"""
Worker Tracker - simple web app (v2).

Flow in the browser:
  1. Upload a video
  2. Draw the ENTRY/EXIT area and the WORK area on the first frame
  3. Click "Start tracking" (live preview, progress, speed)
  4. Get a detailed report for EVERY tracked person: photo, ID, time in each area,
     zone timeline, working / not working verdict - plus downloads
     (HTML/PDF, Excel, CSV, JSON, photos ZIP, processed video)

Only anonymous IDs (Person #1, #2 ...) are used. No face recognition is done.
Person photos are plain crops from the video, saved locally in the jobs/ folder.

Run:
    pip install -r requirements.txt
    python app.py
    -> open http://127.0.0.1:5000
"""

import base64
import csv
import html
import json
import os
import re
import shutil
import subprocess
import threading
import time
import uuid
import zipfile
from collections import defaultdict, deque
from datetime import datetime
from math import hypot
from pathlib import Path

import cv2
import numpy as np
from flask import Flask, abort, jsonify, request, send_from_directory

BASE = Path(__file__).parent
JOBS_DIR = BASE / "jobs"
JOBS_DIR.mkdir(exist_ok=True)

ALLOWED_EXT = {".mp4", ".avi", ".mov", ".mkv", ".wmv", ".m4v", ".webm"}
MODEL_NAME = "yolov8n.pt"          # auto-downloaded on first run
DEBOUNCE_SECONDS = 0.5             # a zone change must last this long to count

BGR = {"ENTRY_EXIT": (0, 200, 255), "WORK_AREA": (0, 200, 0), "OTHER": (200, 200, 200)}
ZONE_HEX = {"ENTRY_EXIT": "#f5a300", "WORK_AREA": "#1fa84f", "OTHER": "#9aa3b2"}
ZONE_LABEL = {"ENTRY_EXIT": "Entry/Exit area", "WORK_AREA": "Work area", "OTHER": "Other area"}
ZONE_SHORT = {"ENTRY_EXIT": "Entry/Exit", "WORK_AREA": "Work", "OTHER": "Other"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024 * 1024   # 1 GB upload limit

JOBS = {}                       # job_id -> status dict (in memory)
RUN_LOCK = threading.Lock()     # process one video at a time


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def job_dir(job_id):
    if len(job_id) != 32 or any(c not in "0123456789abcdef" for c in job_id):
        abort(404)
    d = JOBS_DIR / job_id
    if not d.is_dir():
        abort(404)
    return d


def fmt(sec):
    """Duration, e.g. '2m 05s' or '1h 02m 03s'."""
    sec = int(round(sec))
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}h {m:02d}m {s:02d}s" if h else f"{m}m {s:02d}s"


def fmt_ts(sec):
    """Video timestamp, e.g. '02:05' or '1:02:05'."""
    sec = int(sec)
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def zone_of(point, polys):
    """Zone containing the point. WORK_AREA wins if polygons overlap."""
    for name in ("WORK_AREA", "ENTRY_EXIT"):
        if cv2.pointPolygonTest(polys[name], point, False) >= 0:
            return name
    return "OTHER"


def load_model():
    """Load YOLO, working around the PyTorch >= 2.6 "weights_only" UnpicklingError
    that older ultralytics versions hit (best fix: pip install -U ultralytics).
    Fallback: PyTorch's own TORCH_FORCE_NO_WEIGHTS_ONLY_LOAD switch is turned on for
    this single load only. Acceptable here because the file is the official
    yolov8n.pt that ultralytics downloads itself - don't do this for unknown files."""
    from ultralytics import YOLO
    try:
        return YOLO(MODEL_NAME)
    except Exception as e:                                    # noqa: BLE001
        if "weights_only" not in str(e) and "UnpicklingError" not in repr(e):
            raise
        key = "TORCH_FORCE_NO_WEIGHTS_ONLY_LOAD"
        previous = os.environ.get(key)
        os.environ[key] = "1"
        try:
            return YOLO(MODEL_NAME)
        finally:
            if previous is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = previous


def get_device(preference="auto"):
    """Pick the compute device. Returns (device_string, human_label).
    auto -> NVIDIA GPU (cuda:0) if available, else Apple GPU (mps), else CPU."""
    try:
        import torch
    except Exception:                                         # noqa: BLE001
        return "cpu", "CPU"
    if preference not in ("auto", "cpu") and preference:
        return preference, preference
    if preference != "cpu":
        if torch.cuda.is_available():
            return "cuda:0", f"GPU ({torch.cuda.get_device_name(0)})"
        mps = getattr(torch.backends, "mps", None)
        if mps is not None and mps.is_available():
            return "mps", "Apple GPU (mps)"
    return "cpu", "CPU"


def to_browser_mp4(src, dst):
    """Re-encode to H.264 so browsers can play it. Returns True on success."""
    exe = shutil.which("ffmpeg")
    if not exe:
        try:
            import imageio_ffmpeg
            exe = imageio_ffmpeg.get_ffmpeg_exe()
        except Exception:
            return False
    try:
        subprocess.run(
            [exe, "-y", "-i", str(src), "-vcodec", "libx264", "-pix_fmt", "yuv420p",
             "-movflags", "+faststart", "-loglevel", "error", str(dst)],
            check=True,
        )
        return True
    except Exception:
        return False


# ----------------------------------------------------------------------------
# Per-person tracking data
# ----------------------------------------------------------------------------
class PersonStats:
    """Everything we learn about one tracked ID."""

    def __init__(self, pid, ts, debounce):
        self.pid = pid
        self.first_seen = ts
        self.last_seen = ts
        self.debounce = debounce
        self.zone = None                     # committed (debounced) zone
        self.cand = None                     # zone we may be switching to
        self.cand_since = 0.0
        self.zone_time = defaultdict(float)  # seconds per zone
        self.segments = []                   # [{zone, start, end}] timeline
        self.trail = deque(maxlen=40)
        self.detections = 0
        self.conf_sum = 0.0
        self.distance = 0.0                  # pixels travelled (feet position)
        self.prev_foot = None
        self.best_score = 0.0
        self.best_crop = None

    def update(self, ts, raw_zone, foot, conf, max_gap):
        dt = ts - self.last_seen
        if self.zone is None or dt > max_gap:
            # first sighting, or re-appearance after the track was lost
            self.zone, self.cand = raw_zone, None
            self.segments.append({"zone": raw_zone, "start": ts, "end": ts})
        else:
            if dt > 0:
                self.zone_time[self.zone] += dt
            if raw_zone == self.zone:
                self.cand = None
            elif self.cand != raw_zone:
                self.cand, self.cand_since = raw_zone, ts
            elif ts - self.cand_since >= self.debounce:
                # zone change confirmed: move the waiting time to the new zone
                old, new = self.zone, raw_zone
                moved = min(ts - self.cand_since, self.zone_time[old])
                self.zone_time[old] -= moved
                self.zone_time[new] += moved
                self.segments[-1]["end"] = self.cand_since
                self.segments.append({"zone": new, "start": self.cand_since, "end": ts})
                self.zone, self.cand = new, None
            self.segments[-1]["end"] = ts

        if self.prev_foot is not None and dt <= max_gap:
            self.distance += hypot(foot[0] - self.prev_foot[0], foot[1] - self.prev_foot[1])
        self.prev_foot = foot
        self.last_seen = ts
        self.detections += 1
        self.conf_sum += conf
        self.trail.append(foot)

    def offer_crop(self, frame, box, conf):
        """Keep the best-looking snapshot (big, confident, not cut by the frame edge)."""
        h, w = frame.shape[:2]
        x1, y1, x2, y2 = [float(v) for v in box]
        area = max(x2 - x1, 1.0) * max(y2 - y1, 1.0)
        at_edge = x1 <= 2 or y1 <= 2 or x2 >= w - 2 or y2 >= h - 2
        score = conf * area * (0.5 if at_edge else 1.0)
        if self.best_crop is not None and score <= self.best_score * 1.05:
            return
        px, py = 0.06 * (x2 - x1), 0.04 * (y2 - y1)
        xa, ya = int(max(0, x1 - px)), int(max(0, y1 - py))
        xb, yb = int(min(w, x2 + px)), int(min(h, y2 + py))
        crop = frame[ya:yb, xa:xb]
        if crop.size == 0:
            return
        ch = crop.shape[0]
        if ch > 260:
            crop = cv2.resize(crop, (max(1, int(crop.shape[1] * 260 / ch)), 260),
                              interpolation=cv2.INTER_AREA)
        self.best_score, self.best_crop = score, crop.copy()

    @property
    def total(self):
        return self.last_seen - self.first_seen


def summarize_person(p, video_duration, opts):
    z = p.zone_time
    work, entry, other = z["WORK_AREA"], z["ENTRY_EXIT"], z["OTHER"]
    counted = work + entry + other
    ratio = work / counted if counted > 0 else 0.0
    brief = p.total < opts["min_seconds"]
    if brief:
        verdict = "BRIEF"
    else:
        verdict = "WORKING" if ratio >= opts["work_threshold"] else "NOT WORKING"

    segs = []
    for s in p.segments:
        dur = max(s["end"] - s["start"], 0.0)
        segs.append({
            "zone": s["zone"], "zone_label": ZONE_LABEL[s["zone"]],
            "start": round(s["start"], 2), "end": round(s["end"], 2),
            "duration": round(dur, 2),
            "start_fmt": fmt_ts(s["start"]), "end_fmt": fmt_ts(s["end"]),
            "duration_fmt": fmt(dur),
        })

    path = []
    for s in segs:
        label = ZONE_SHORT[s["zone"]]
        if not path or path[-1] != label:
            path.append(label)
    work_durs = [s["duration"] for s in segs if s["zone"] == "WORK_AREA"]

    first_zone, last_zone = segs[0]["zone"], segs[-1]["zone"]
    if p.first_seen < 1.0:
        entered_via = "Already in frame at start of video"
    elif first_zone == "ENTRY_EXIT":
        entered_via = "Via Entry/Exit area"
    else:
        entered_via = f"Appeared in {ZONE_LABEL[first_zone]}"
    if p.last_seen >= video_duration - 1.0:
        ended = "Still in frame at end of video"
    elif last_zone == "ENTRY_EXIT":
        ended = "Left via Entry/Exit area"
    else:
        ended = f"Lost from view in {ZONE_LABEL[last_zone]}"

    return {
        "id": p.pid, "name": f"Person #{p.pid}", "verdict": verdict, "brief": brief,
        "first_seen_s": round(p.first_seen, 2), "last_seen_s": round(p.last_seen, 2),
        "first_seen": fmt_ts(p.first_seen), "last_seen": fmt_ts(p.last_seen),
        "total_s": round(p.total, 2), "total_time": fmt(p.total),
        "work_s": round(work, 2), "work_time": fmt(work),
        "entry_exit_s": round(entry, 2), "entry_exit_time": fmt(entry),
        "other_s": round(other, 2), "other_time": fmt(other),
        "work_percent": round(ratio * 100, 1),
        "work_sessions": sum(1 for lab in path if lab == "Work"),
        "longest_work_s": round(max(work_durs), 2) if work_durs else 0.0,
        "longest_work": fmt(max(work_durs)) if work_durs else "0m 00s",
        "zone_changes": max(len(path) - 1, 0),
        "zone_path": " > ".join(path),
        "entered_via": entered_via, "ended": ended,
        "distance_px": int(round(p.distance)),
        "detections": p.detections,
        "avg_confidence": round(p.conf_sum / p.detections, 2) if p.detections else 0.0,
        "segments": segs,
        "photo": None, "photo_w": 0, "photo_h": 0,
    }


# ----------------------------------------------------------------------------
# Report writers
# ----------------------------------------------------------------------------
PEOPLE_COLUMNS = [
    ("name", "Person"), ("verdict", "Verdict"), ("first_seen", "First seen"),
    ("last_seen", "Last seen"), ("total_time", "Time on screen"),
    ("work_time", "Time in work area"), ("entry_exit_time", "Time in entry/exit area"),
    ("other_time", "Time elsewhere"), ("work_percent", "Work %"),
    ("work_sessions", "Work sessions"), ("longest_work", "Longest work stretch"),
    ("zone_changes", "Zone changes"), ("entered_via", "How they appeared"),
    ("ended", "How they ended"), ("zone_path", "Zone path"),
    ("distance_px", "Movement (px)"), ("detections", "Detections"),
    ("avg_confidence", "Avg detection confidence"),
]


def write_csv(d, people):
    with open(d / "report.csv", "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.writer(f)
        wr.writerow([label for _, label in PEOPLE_COLUMNS] + ["Photo file"])
        for p in people:
            wr.writerow([p[key] for key, _ in PEOPLE_COLUMNS] + [p["photo"] or ""])
    with open(d / "zone_timeline.csv", "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.writer(f)
        wr.writerow(["Person", "Zone", "Start", "End", "Duration (s)", "Duration"])
        for p in people:
            for s in p["segments"]:
                wr.writerow([p["name"], s["zone_label"], s["start_fmt"], s["end_fmt"],
                             s["duration"], s["duration_fmt"]])


def write_json(d, result):
    with open(d / "report.json", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2)


def write_zip(d, people):
    photos = [p for p in people if p["photo"]]
    if not photos:
        return False
    with zipfile.ZipFile(d / "people.zip", "w", zipfile.ZIP_STORED) as z:
        for p in photos:
            z.write(d / p["photo"], f"person_{p['id']}.jpg")
    return True


def write_xlsx(d, result):
    """Excel report with photos embedded. Returns False if openpyxl/Pillow is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:                                         # noqa: BLE001
        return False

    v, s, st, people = result["video"], result["summary"], result["settings"], result["people"]
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="2F5BEA")
    head_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Video", v["name"]), ("Duration", v["duration"]),
        ("Resolution", f'{v["width"]} x {v["height"]} @ {v["fps"]} fps'),
        ("Analysed", v["analysed_at"]), ("Processing time", fmt(v["processing_s"])),
        ("Device", v["device"]), ("Tracker", st["tracker"]),
        ("Working if time in work area >=", f'{st["work_threshold_percent"]}%'),
        ("Brief detection if seen less than", f'{st["brief_seconds"]} s'),
        ("", ""),
        ("People tracked (IDs)", s["tracked_ids"]), ("Judged (long enough)", s["judged"]),
        ("Working", s["working"]), ("Not working", s["not_working"]),
        ("Brief detections", s["brief"]), ("Average work % (judged)", s["avg_work_percent"]),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 40

    ws2 = wb.create_sheet("People")
    ws2.append(["Photo"] + [label for _, label in PEOPLE_COLUMNS])
    for c in ws2[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, p in enumerate(people, start=2):
        ws2.append([""] + [p[key] for key, _ in PEOPLE_COLUMNS])
        ws2.row_dimensions[i].height = 80
        for c in ws2[i]:
            c.alignment = Alignment(vertical="center", wrap_text=True)
        if p["photo"]:
            try:
                img = XLImage(str(d / p["photo"]))
                ratio = p["photo_w"] / max(p["photo_h"], 1)
                img.height, img.width = 100, int(100 * ratio)
                ws2.add_image(img, f"A{i}")
            except Exception:                                 # noqa: BLE001
                pass
    ws2.column_dimensions["A"].width = 16
    for idx in range(2, len(PEOPLE_COLUMNS) + 2):
        ws2.column_dimensions[get_column_letter(idx)].width = 18
    ws2.column_dimensions["O"].width = 30
    ws2.column_dimensions["P"].width = 34
    ws2.freeze_panes = "B2"

    ws3 = wb.create_sheet("Zone timeline")
    ws3.append(["Person", "Zone", "Start", "End", "Duration (s)", "Duration"])
    for c in ws3[1]:
        c.fill, c.font = head_fill, head_font
    for p in people:
        for sg in p["segments"]:
            ws3.append([p["name"], sg["zone_label"], sg["start_fmt"], sg["end_fmt"],
                        sg["duration"], sg["duration_fmt"]])
    for col, w in zip("ABCDEF", (14, 18, 10, 10, 14, 14)):
        ws3.column_dimensions[col].width = w

    wb.save(d / "report.xlsx")
    return True


REPORT_CSS = """
*{box-sizing:border-box} body{margin:0;font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif;
background:#f5f6f8;color:#1c2230} .wrap{max-width:960px;margin:0 auto;padding:24px 16px}
h1{margin:0 0 4px;font-size:24px} .sub{color:#6b7385;font-size:13px;margin-bottom:18px}
.info{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:8px;margin-bottom:16px}
.info div{background:#fff;border:1px solid #e3e6ec;border-radius:8px;padding:8px 12px}
.info span{display:block;color:#6b7385;font-size:11px;text-transform:uppercase}
.tiles{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:16px}
.tile{background:#fff;border:1px solid #e3e6ec;border-radius:8px;padding:10px 18px}
.tile b{display:block;font-size:24px}.tile span{color:#6b7385;font-size:12px}
.legend{font-size:12px;color:#6b7385;margin:0 0 12px}.legend i{display:inline-block;width:11px;height:11px;
border-radius:2px;margin:0 4px 0 12px;vertical-align:-1px}
.card{display:flex;gap:16px;background:#fff;border:1px solid #e3e6ec;border-radius:12px;padding:14px;
margin-bottom:12px;break-inside:avoid}.card.brief{opacity:.75}
.ph{width:110px;flex:none}.ph img{width:110px;height:150px;object-fit:cover;border-radius:8px;background:#000}
.noph{width:110px;height:150px;border-radius:8px;background:#e3e6ec;color:#6b7385;display:flex;
align-items:center;justify-content:center;font-size:12px}
.body{flex:1;min-width:0}.head{display:flex;gap:10px;align-items:center;margin-bottom:8px}
.head h3{margin:0;font-size:18px}
.badge{padding:3px 10px;border-radius:20px;font-size:12px;font-weight:600;color:#fff}
.w{background:#1fa84f}.n{background:#d9534f}.b{background:#9aa3b2}
.kv{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:6px 14px;font-size:13px}
.kv span{display:block;color:#6b7385;font-size:11px;text-transform:uppercase}
.tl{position:relative;height:14px;background:#eceef3;border-radius:7px;margin:12px 0 2px;overflow:hidden}
.tl span{position:absolute;top:0;bottom:0}
.scale{display:flex;justify-content:space-between;font-size:11px;color:#6b7385}
.note{color:#6b7385;font-size:12px;margin-top:20px}
button{font:inherit;padding:8px 14px;border-radius:8px;border:1px solid #2f5bea;background:#2f5bea;
color:#fff;cursor:pointer;margin-bottom:12px}
@media print{body{background:#fff}.noprint{display:none}.card,.tile,.info div{border-color:#ccc}}
"""


def timeline_html(p, duration):
    dur = max(duration, 0.001)
    parts = []
    for s in p["segments"]:
        left = s["start"] / dur * 100
        width = max((s["end"] - s["start"]) / dur * 100, 0.4)
        title = html.escape(f'{s["zone_label"]} {s["start_fmt"]}-{s["end_fmt"]}')
        parts.append(f'<span style="left:{left:.2f}%;width:{width:.2f}%;'
                     f'background:{ZONE_HEX[s["zone"]]}" title="{title}"></span>')
    return '<div class="tl">' + "".join(parts) + "</div>"


def build_html_report(d, result):
    esc = html.escape
    v, s, st, people = result["video"], result["summary"], result["settings"], result["people"]
    info = [
        ("Video", v["name"]), ("Duration", v["duration"]),
        ("Resolution", f'{v["width"]} x {v["height"]} @ {v["fps"]} fps'),
        ("Analysed", v["analysed_at"]), ("Device", v["device"]), ("Tracker", st["tracker"]),
        ("Working rule", f'at least {st["work_threshold_percent"]}% of time in work area'),
        ("Brief detection", f'seen less than {st["brief_seconds"]} s'),
    ]
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>Worker report - {esc(v['name'])}</title><style>{REPORT_CSS}</style></head>"
           f"<body><div class='wrap'><h1>Worker Tracking Report</h1>"
           f"<div class='sub'>Anonymous tracking - temporary IDs only, no face recognition.</div>"
           f"<button class='noprint' onclick='window.print()'>Print / Save as PDF</button>"]
    out.append("<div class='info'>" + "".join(
        f"<div><span>{esc(k)}</span>{esc(str(val))}</div>" for k, val in info) + "</div>")
    out.append(
        f"<div class='tiles'><div class='tile'><b>{s['tracked_ids']}</b><span>IDs tracked</span></div>"
        f"<div class='tile'><b>{s['working']}</b><span>Working</span></div>"
        f"<div class='tile'><b>{s['not_working']}</b><span>Not working</span></div>"
        f"<div class='tile'><b>{s['brief']}</b><span>Brief detections</span></div>"
        f"<div class='tile'><b>{s['avg_work_percent']}%</b><span>Avg work time (judged)</span></div></div>")
    out.append("<p class='legend'>Timeline colours:"
               + "".join(f"<i style='background:{ZONE_HEX[z]}'></i>{ZONE_LABEL[z]}"
                         for z in ("ENTRY_EXIT", "WORK_AREA", "OTHER"))
               + "<br>Each bar spans the whole video; coloured parts show where the person was.</p>")

    for p in people:
        if p["photo"]:
            b64 = base64.b64encode((d / p["photo"]).read_bytes()).decode()
            photo = f"<img src='data:image/jpeg;base64,{b64}' alt='{esc(p['name'])}'>"
        else:
            photo = "<div class='noph'>No photo</div>"
        cls = {"WORKING": "w", "NOT WORKING": "n"}.get(p["verdict"], "b")
        note = (f"<div style='font-size:12px;color:#6b7385;margin-bottom:6px'>Seen for only "
                f"{p['total_s']:.1f} s - too short to judge (may be a false detection or a split ID).</div>"
                if p["brief"] else "")
        kv = [
            ("Seen", f"{p['first_seen']} - {p['last_seen']}"), ("On screen", p["total_time"]),
            ("In work area", f"{p['work_time']} ({p['work_percent']}%)"),
            ("In entry/exit area", p["entry_exit_time"]), ("Elsewhere", p["other_time"]),
            ("Work sessions", p["work_sessions"]), ("Longest work stretch", p["longest_work"]),
            ("Zone changes", p["zone_changes"]), ("How they appeared", p["entered_via"]),
            ("How they ended", p["ended"]), ("Zone path", p["zone_path"]),
            ("Movement", f"{p['distance_px']} px"),
            ("Detection confidence", f"{p['avg_confidence']} ({p['detections']} frames)"),
        ]
        kvh = "".join(f"<div><span>{esc(k)}</span>{esc(str(val))}</div>" for k, val in kv)
        out.append(
            f"<div class='card {'brief' if p['brief'] else ''}'><div class='ph'>{photo}</div>"
            f"<div class='body'><div class='head'><h3>{esc(p['name'])}</h3>"
            f"<span class='badge {cls}'>{esc(p['verdict'])}</span></div>{note}"
            f"<div class='kv'>{kvh}</div>{timeline_html(p, v['duration_s'])}"
            f"<div class='scale'><span>00:00</span><span>{esc(fmt_ts(v['duration_s']))}</span></div>"
            f"</div></div>")
    if not people:
        out.append("<p>No people were tracked. Check the zones and the video.</p>")
    out.append("<p class='note'>Note: a person who leaves view and comes back may receive a new ID, "
               "so one real person can appear as more than one entry. 'Working' means being inside "
               "the marked work area, not a measurement of activity.</p></div></body></html>")
    (d / "report.html").write_text("".join(out), encoding="utf-8")


# ----------------------------------------------------------------------------
# Background processing
# ----------------------------------------------------------------------------
def process_job(job_id, zones_raw, opts):
    job = JOBS[job_id]
    d = JOBS_DIR / job_id
    job.update(state="queued", progress=0, message="Waiting for the tracker...")

    try:
        with RUN_LOCK:
            t_start = time.time()
            job.update(state="running", message="Loading model...")
            model = load_model()              # fresh model => fresh tracker state
            device, device_label = get_device()
            use_half = device.startswith("cuda")          # fp16 = faster on NVIDIA GPUs
            print(f"[job {job_id[:8]}] running on {device_label}")

            polys = {k: np.array(v, np.int32) for k, v in zones_raw.items()}
            cap = cv2.VideoCapture(str(d / job["video"]))
            fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) or 0
            w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
            h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
            raw_path = d / "raw.mp4"
            writer = cv2.VideoWriter(str(raw_path), cv2.VideoWriter_fourcc(*"mp4v"),
                                     fps, (w, h))

            people = {}
            frame_idx = 0
            t_loop = time.time()
            job["message"] = f"Tracking people on {device_label}"

            while True:
                ok, raw = cap.read()
                if not ok:
                    break
                ts = frame_idx / fps

                res = model.track(raw, persist=True, classes=[0], conf=opts["conf"],
                                  tracker=opts["tracker"], device=device, half=use_half,
                                  verbose=False)[0]

                # draw zones on a copy (raw stays clean for person snapshots)
                overlay = raw.copy()
                for name, poly in polys.items():
                    cv2.fillPoly(overlay, [poly], BGR[name])
                frame = cv2.addWeighted(overlay, 0.2, raw, 0.8, 0)
                for name, poly in polys.items():
                    cv2.polylines(frame, [poly], True, BGR[name], 2)
                    cv2.putText(frame, name, tuple(int(v) for v in poly[0]),
                                cv2.FONT_HERSHEY_SIMPLEX, 0.6, BGR[name], 2)

                in_frame = 0
                if res.boxes is not None and res.boxes.id is not None:
                    ids = res.boxes.id.int().cpu().tolist()
                    boxes = res.boxes.xyxy.cpu().numpy()
                    confs = res.boxes.conf.cpu().numpy()
                    in_frame = len(ids)
                    for pid, box, conf in zip(ids, boxes, confs):
                        x1, y1, x2, y2 = box
                        foot = (int((x1 + x2) / 2), int(y2))     # feet position
                        zone = zone_of(foot, polys)
                        if pid not in people:
                            people[pid] = PersonStats(pid, ts, DEBOUNCE_SECONDS)
                        p = people[pid]
                        p.update(ts, zone, foot, float(conf), opts["max_gap"])
                        if opts["snapshots"]:
                            p.offer_crop(raw, box, float(conf))

                        c = BGR[zone]
                        cv2.rectangle(frame, (int(x1), int(y1)), (int(x2), int(y2)), c, 2)
                        cv2.putText(frame, f"Person #{pid} [{zone}]",
                                    (int(x1), max(int(y1) - 6, 12)),
                                    cv2.FONT_HERSHEY_SIMPLEX, 0.55, c, 2)
                        pts = np.array(p.trail, np.int32)
                        if len(pts) > 1:
                            cv2.polylines(frame, [pts], False, c, 2)

                writer.write(frame)
                frame_idx += 1

                # live preview + progress
                job["people_seen"] = len(people)
                job["people_now"] = in_frame
                if frame_idx % 8 == 1:
                    tmp = d / "preview_tmp.jpg"
                    cv2.imwrite(str(tmp), frame)
                    os.replace(tmp, d / "preview.jpg")
                    job["preview_ts"] = frame_idx
                    speed = frame_idx / max(time.time() - t_loop, 1e-6)
                    msg = f"Tracking people on {device_label} - {speed:.1f} frames/s"
                    if total_frames:
                        job["progress"] = min(99, int(frame_idx / total_frames * 100))
                        msg += f" - about {fmt(max(total_frames - frame_idx, 0) / speed)} left"
                    job["message"] = msg

            cap.release()
            writer.release()
            video_duration = frame_idx / fps
            processing_s = time.time() - t_start

        # ---- encode for browser (outside the lock: doesn't need the model) ----
        job["message"] = "Encoding video..."
        playable = to_browser_mp4(raw_path, d / "processed.mp4")
        if not playable:
            shutil.copy(raw_path, d / "processed.mp4")
        job["video_playable"] = playable

        # ---- build the report ----
        job["message"] = "Building report..."
        (d / "people").mkdir(exist_ok=True)
        rows = []
        for pid in sorted(people):
            p = people[pid]
            row = summarize_person(p, video_duration, opts)
            if p.best_crop is not None:
                rel = f"people/person_{pid}.jpg"
                cv2.imwrite(str(d / rel), p.best_crop, [cv2.IMWRITE_JPEG_QUALITY, 90])
                row["photo"] = rel
                row["photo_h"], row["photo_w"] = p.best_crop.shape[:2]
            rows.append(row)

        judged = [r for r in rows if not r["brief"]]
        result = {
            "video": {
                "name": job.get("original_name", "video"),
                "duration_s": round(video_duration, 2), "duration": fmt(video_duration),
                "width": w, "height": h, "fps": round(fps, 2), "frames": frame_idx,
                "analysed_at": datetime.now().strftime("%d %b %Y, %H:%M"),
                "processing_s": round(processing_s, 1), "device": device_label,
            },
            "settings": {
                "work_threshold_percent": round(opts["work_threshold"] * 100),
                "brief_seconds": opts["min_seconds"], "tracker": opts["tracker"],
                "confidence": opts["conf"], "zone_debounce_s": DEBOUNCE_SECONDS,
                "zones": zones_raw,
            },
            "summary": {
                "tracked_ids": len(rows), "judged": len(judged),
                "working": sum(1 for r in judged if r["verdict"] == "WORKING"),
                "not_working": sum(1 for r in judged if r["verdict"] == "NOT WORKING"),
                "brief": len(rows) - len(judged),
                "avg_work_percent": round(sum(r["work_percent"] for r in judged) / len(judged), 1)
                                    if judged else 0.0,
            },
            "people": rows,
        }

        write_csv(d, rows)
        write_json(d, result)
        build_html_report(d, result)
        has_xlsx = write_xlsx(d, result)
        has_zip = write_zip(d, rows)

        downloads = [
            {"label": "Report - view / save as PDF", "file": "report.html", "newtab": True,
             "note": "opens in a new tab, then Print > Save as PDF"},
            {"label": "Report - HTML file", "file": "report.html", "note": "photos included"},
        ]
        if has_xlsx:
            downloads.append({"label": "Excel (.xlsx)", "file": "report.xlsx",
                              "note": "photos + 3 sheets"})
        downloads += [
            {"label": "CSV - people summary", "file": "report.csv", "note": "one row per person"},
            {"label": "CSV - zone timeline", "file": "zone_timeline.csv",
             "note": "every zone stay"},
            {"label": "JSON - full detail", "file": "report.json", "note": "for developers"},
        ]
        if has_zip:
            downloads.append({"label": "Person photos (.zip)", "file": "people.zip",
                              "note": "one image per ID"})
        downloads.append({"label": "Processed video (.mp4)", "file": "processed.mp4",
                          "note": "boxes, IDs, zones, trails"})

        job.update(state="done", progress=100, message="Done", result=result,
                   downloads=downloads)

    except Exception as e:                                     # noqa: BLE001
        job.update(state="error", message=f"{type(e).__name__}: {e}", error=True)


# ----------------------------------------------------------------------------
# Routes
# ----------------------------------------------------------------------------
@app.get("/")
def index():
    return INDEX_HTML


@app.post("/upload")
def upload():
    f = request.files.get("video")
    if not f or not f.filename:
        return jsonify(error="No file received"), 400
    ext = Path(f.filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        return jsonify(error=f"Unsupported file type {ext}"), 400

    job_id = uuid.uuid4().hex
    d = JOBS_DIR / job_id
    d.mkdir()
    video_name = f"input{ext}"
    f.save(d / video_name)

    cap = cv2.VideoCapture(str(d / video_name))
    ok, frame = cap.read()
    cap.release()
    if not ok:
        shutil.rmtree(d, ignore_errors=True)
        return jsonify(error="Could not read that video"), 400
    cv2.imwrite(str(d / "frame.jpg"), frame)

    JOBS[job_id] = {"state": "uploaded", "video": video_name, "progress": 0,
                    "original_name": f.filename}
    return jsonify(job_id=job_id, width=frame.shape[1], height=frame.shape[0],
                   name=f.filename)


@app.post("/start/<job_id>")
def start(job_id):
    job_dir(job_id)
    if job_id not in JOBS:
        abort(404)
    if JOBS[job_id]["state"] in ("queued", "running"):
        return jsonify(error="Already running"), 409

    data = request.get_json(force=True, silent=True) or {}
    zones = data.get("zones", {})
    clean = {}
    for name in ("ENTRY_EXIT", "WORK_AREA"):
        pts = zones.get(name, [])
        if len(pts) < 3 or any(len(p) != 2 for p in pts):
            return jsonify(error=f"Zone {name} needs at least 3 points"), 400
        clean[name] = [[int(round(float(x))), int(round(float(y)))] for x, y in pts]

    opts = {
        "work_threshold": min(max(float(data.get("work_threshold", 0.5)), 0), 1),
        "min_seconds": max(float(data.get("min_seconds", 2)), 0),
        "tracker": data.get("tracker") if data.get("tracker") in
                   ("bytetrack.yaml", "botsort.yaml") else "bytetrack.yaml",
        "snapshots": bool(data.get("snapshots", True)),
        "conf": 0.4,
        "max_gap": 1.0,
    }
    JOBS[job_id].update(state="queued", result=None, downloads=None, error=False)
    threading.Thread(target=process_job, args=(job_id, clean, opts), daemon=True).start()
    return jsonify(ok=True)


@app.get("/status/<job_id>")
def status(job_id):
    job_dir(job_id)
    job = JOBS.get(job_id)
    if not job:
        abort(404)
    keys = ("state", "progress", "message", "result", "downloads", "video_playable",
            "people_seen", "people_now", "preview_ts")
    return jsonify({k: job.get(k) for k in keys})


FILE_PREFIX = {".mp4": "processed_", ".zip": "people_"}
STATIC_FILES = {"frame.jpg", "preview.jpg", "processed.mp4", "report.csv", "zone_timeline.csv",
                "report.json", "report.html", "report.xlsx", "people.zip"}


@app.get("/jobs/<job_id>/<path:filename>")
def job_file(job_id, filename):
    d = job_dir(job_id)
    if filename not in STATIC_FILES and not re.fullmatch(r"people/person_\d+\.jpg", filename):
        abort(404)
    if request.args.get("dl"):
        stem = re.sub(r"[^A-Za-z0-9_-]+", "_",
                      Path(JOBS.get(job_id, {}).get("original_name", "video")).stem)[:40] or "video"
        ext = Path(filename).suffix
        base = FILE_PREFIX.get(ext, "worker_report_")
        if filename == "zone_timeline.csv":
            base = "zone_timeline_"
        return send_from_directory(d, filename, as_attachment=True, max_age=0,
                                   download_name=f"{base}{stem}{ext}")
    return send_from_directory(d, filename, max_age=0)


# ----------------------------------------------------------------------------
# Frontend (single page)
# ----------------------------------------------------------------------------
INDEX_HTML = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Worker Tracker</title>
<style>
  :root { --bg:#f5f6f8; --card:#fff; --ink:#1c2230; --muted:#6b7385; --line:#e3e6ec;
          --accent:#2f5bea; --ok:#1fa84f; --warn:#d9534f; --entry:#f5a300; --other:#9aa3b2; }
  * { box-sizing: border-box; }
  body { margin:0; font-family: system-ui, -apple-system, Segoe UI, Roboto, sans-serif;
         background:var(--bg); color:var(--ink); }
  header { padding:18px 24px; background:var(--card); border-bottom:1px solid var(--line); }
  header h1 { margin:0; font-size:20px; }
  header p { margin:4px 0 0; color:var(--muted); font-size:13px; }
  main { max-width:980px; margin:24px auto; padding:0 16px; display:grid; gap:16px; }
  section { background:var(--card); border:1px solid var(--line); border-radius:12px; padding:18px; }
  section h2 { margin:0 0 12px; font-size:16px; display:flex; align-items:center; gap:8px; }
  section h3 { margin:18px 0 8px; font-size:14px; }
  .num { background:var(--accent); color:#fff; width:22px; height:22px; border-radius:50%;
         display:inline-flex; align-items:center; justify-content:center; font-size:12px; }
  .hidden { display:none !important; }
  button, .btn { font:inherit; padding:8px 14px; border-radius:8px; border:1px solid var(--line);
           background:#fff; cursor:pointer; color:var(--ink); text-decoration:none; display:inline-block; }
  button:hover, .btn:hover { border-color:var(--accent); }
  button.primary { background:var(--accent); color:#fff; border-color:var(--accent); }
  button:disabled { opacity:.45; cursor:not-allowed; }
  .row { display:flex; flex-wrap:wrap; gap:8px; align-items:center; margin-bottom:12px; }
  .zbtn.active { outline:2px solid var(--accent); }
  .dot { display:inline-block; width:10px; height:10px; border-radius:50%; margin-right:6px; }
  canvas, .frame { width:100%; height:auto; border-radius:8px; border:1px solid var(--line);
                   display:block; background:#000; }
  canvas { cursor:crosshair; }
  .hint { color:var(--muted); font-size:13px; margin:8px 0 0; }
  .bar { height:10px; background:var(--line); border-radius:6px; overflow:hidden; margin:10px 0; }
  .bar > div { height:100%; width:0; background:var(--accent); transition:width .3s; }
  video { width:100%; border-radius:8px; background:#000; margin-top:14px; }
  details { margin-top:6px; } summary { cursor:pointer; color:var(--muted); font-size:13px; }
  .adv { display:grid; grid-template-columns:repeat(auto-fit,minmax(200px,1fr)); gap:12px; margin-top:10px; }
  label { font-size:13px; color:var(--muted); display:block; }
  input[type=number], select { width:100%; padding:6px 8px; margin-top:4px; border:1px solid var(--line); border-radius:6px; font:inherit; }
  .chk { display:flex; gap:8px; align-items:center; margin-top:22px; }
  .err { color:var(--warn); font-size:14px; margin-top:8px; }

  .info { display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:8px; margin-bottom:14px; }
  .info div { background:var(--bg); border-radius:8px; padding:8px 12px; font-size:13px; word-break:break-word; }
  .info span { display:block; color:var(--muted); font-size:11px; text-transform:uppercase; }
  .stats { display:flex; gap:10px; flex-wrap:wrap; margin-bottom:14px; }
  .stat { background:var(--bg); border-radius:8px; padding:10px 18px; }
  .stat b { display:block; font-size:22px; }
  .stat span { color:var(--muted); font-size:12px; }
  .legend { font-size:12px; color:var(--muted); margin:0 0 10px; }
  .legend i { display:inline-block; width:11px; height:11px; border-radius:2px; margin:0 4px 0 12px; vertical-align:-1px; }

  .card { display:flex; gap:16px; border:1px solid var(--line); border-radius:12px; padding:14px; margin-bottom:12px; }
  .card.brief { opacity:.78; }
  .ph { width:110px; flex:none; }
  .ph img { width:110px; height:150px; object-fit:cover; border-radius:8px; background:#000; display:block; }
  .noph { width:110px; height:150px; border-radius:8px; background:var(--line); color:var(--muted);
          display:flex; align-items:center; justify-content:center; font-size:12px; }
  .cbody { flex:1; min-width:0; }
  .chead { display:flex; gap:10px; align-items:center; margin-bottom:8px; flex-wrap:wrap; }
  .chead h4 { margin:0; font-size:17px; }
  .badge { padding:3px 10px; border-radius:20px; font-size:12px; font-weight:600; color:#fff; }
  .badge.w { background:var(--ok); } .badge.n { background:var(--warn); } .badge.b { background:var(--other); }
  .kv { display:grid; grid-template-columns:repeat(auto-fit,minmax(165px,1fr)); gap:6px 14px; font-size:13px; }
  .kv span { display:block; color:var(--muted); font-size:11px; text-transform:uppercase; }
  .tl { position:relative; height:14px; background:#eceef3; border-radius:7px; margin:12px 0 2px; overflow:hidden; }
  .tl span { position:absolute; top:0; bottom:0; }
  .scale { display:flex; justify-content:space-between; font-size:11px; color:var(--muted); }
  .seg { width:100%; border-collapse:collapse; font-size:12px; margin-top:6px; }
  .seg th, .seg td { text-align:left; padding:3px 8px; border-bottom:1px solid var(--line); }
  .dl { display:grid; grid-template-columns:repeat(auto-fit,minmax(210px,1fr)); gap:8px; }
  .dl .btn { text-align:left; } .dl small { display:block; color:var(--muted); font-size:11px; margin-top:2px; }
  @media (max-width:560px) { .card { flex-direction:column; } }
</style>
</head>
<body>
<header>
  <h1>Worker Tracker</h1>
  <p>Anonymous people tracking &mdash; upload a video, mark the areas, get a detailed work report.</p>
</header>

<main>
  <!-- STEP 1 -->
  <section id="s1">
    <h2><span class="num">1</span> Select video</h2>
    <div class="row">
      <input type="file" id="file" accept="video/*">
      <button class="primary" id="upBtn">Upload</button>
    </div>
    <div class="hint" id="upMsg">Supported: mp4, avi, mov, mkv, webm, wmv, m4v</div>
    <div class="err" id="upErr"></div>
  </section>

  <!-- STEP 2 -->
  <section id="s2" class="hidden">
    <h2><span class="num">2</span> Mark the areas on the first frame</h2>
    <div class="row">
      <button class="zbtn active" id="zEntry"><span class="dot" style="background:var(--entry)"></span>Entry / Exit area</button>
      <button class="zbtn" id="zWork"><span class="dot" style="background:var(--ok)"></span>Work area</button>
      <button id="undo">Undo point</button>
      <button id="clear">Clear this area</button>
    </div>
    <canvas id="cv"></canvas>
    <p class="hint">Pick an area above, then click its corners on the image (at least 3 points).
       Draw the work area on the floor where people stand while working &mdash; the app uses their feet position.</p>

    <details>
      <summary>Advanced settings</summary>
      <div class="adv">
        <label>Working if time in work area is at least (%)
          <input type="number" id="thr" value="50" min="0" max="100"></label>
        <label>Mark people seen less than (seconds) as "brief"
          <input type="number" id="minsec" value="2" min="0"></label>
        <label>Tracker
          <select id="trk"><option value="bytetrack.yaml">ByteTrack (fast)</option>
                           <option value="botsort.yaml">BoT-SORT (better with overlap)</option></select></label>
        <label class="chk"><input type="checkbox" id="snap" checked> Save a photo (video crop) of each person</label>
      </div>
    </details>

    <div class="row" style="margin-top:14px">
      <button class="primary" id="go" disabled>Start tracking</button>
      <span class="hint" id="goHint">Draw both areas to continue.</span>
    </div>
    <div class="err" id="goErr"></div>
  </section>

  <!-- STEP 3 -->
  <section id="s3" class="hidden">
    <h2><span class="num">3</span> Processing</h2>
    <div class="hint" id="msg">Starting...</div>
    <div class="bar"><div id="pb"></div></div>
    <div class="hint" id="live"></div>
    <img id="prev" class="frame hidden" alt="live preview" style="margin-top:10px">
  </section>

  <!-- RESULTS -->
  <section id="s4" class="hidden">
    <h2><span class="num">&#10003;</span> Results</h2>
    <div class="info" id="vinfo"></div>
    <div class="stats" id="tiles"></div>

    <div class="row">
      <label style="display:flex;gap:8px;align-items:center">Show
        <select id="flt" style="width:auto;margin:0">
          <option value="all">everyone</option>
          <option value="WORKING">working</option>
          <option value="NOT WORKING">not working</option>
          <option value="BRIEF">brief detections</option>
        </select></label>
    </div>
    <p class="legend">Timeline:
      <i style="background:var(--entry)"></i>Entry/Exit area
      <i style="background:var(--ok)"></i>Work area
      <i style="background:var(--other)"></i>Elsewhere
      &nbsp;&middot; each bar spans the whole video</p>
    <div id="people"></div>
    <p class="hint" id="noRows" style="display:none">No one was tracked. Check the zones, or lower the confidence of the video quality.</p>
    <p class="hint">A person who leaves view and returns may get a new ID, so one real person can appear more than once.
       &ldquo;Working&rdquo; means being inside the work area, not a measure of activity.</p>

    <h3>Download report</h3>
    <div class="dl" id="dl"></div>

    <video id="vid" controls class="hidden"></video>
    <p class="hint hidden" id="noPlay">The video could not be converted for in-browser playback (ffmpeg missing) &mdash; use the download button.</p>
    <div class="row" style="margin-top:14px"><button id="again">Process another video</button></div>
  </section>
</main>

<script>
const $ = id => document.getElementById(id);
const esc = s => String(s).replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const COLORS = { ENTRY_EXIT: "#f5a300", WORK_AREA: "#1fa84f" };
const LABELS = { ENTRY_EXIT: "ENTRY / EXIT", WORK_AREA: "WORK AREA" };
const ZH = { ENTRY_EXIT: "#f5a300", WORK_AREA: "#1fa84f", OTHER: "#9aa3b2" };
let jobId = null, img = new Image(), RES = null;
let zones = { ENTRY_EXIT: [], WORK_AREA: [] };
let active = "ENTRY_EXIT";
let timer = null;

/* ---------- Step 1: upload ---------- */
$("upBtn").onclick = async () => {
  const f = $("file").files[0];
  $("upErr").textContent = "";
  if (!f) { $("upErr").textContent = "Choose a video file first."; return; }
  const fd = new FormData(); fd.append("video", f);
  $("upBtn").disabled = true; $("upMsg").textContent = "Uploading...";
  try {
    const r = await fetch("/upload", { method: "POST", body: fd });
    const j = await r.json();
    if (!r.ok) throw new Error(j.error || "Upload failed");
    jobId = j.job_id;
    $("upMsg").textContent = "Uploaded: " + j.name;
    img.onload = () => { initCanvas(); $("s2").classList.remove("hidden");
                         $("s2").scrollIntoView({behavior:"smooth"}); };
    img.src = `/jobs/${jobId}/frame.jpg?t=${Date.now()}`;
  } catch (e) { $("upErr").textContent = e.message; $("upMsg").textContent = ""; }
  $("upBtn").disabled = false;
};

/* ---------- Step 2: draw zones ---------- */
const cv = $("cv"), ctx = cv.getContext("2d");
function initCanvas() {
  cv.width = img.naturalWidth; cv.height = img.naturalHeight;
  zones = { ENTRY_EXIT: [], WORK_AREA: [] }; setActive("ENTRY_EXIT"); redraw();
}
function setActive(z) {
  active = z;
  $("zEntry").classList.toggle("active", z === "ENTRY_EXIT");
  $("zWork").classList.toggle("active", z === "WORK_AREA");
}
$("zEntry").onclick = () => setActive("ENTRY_EXIT");
$("zWork").onclick = () => setActive("WORK_AREA");
$("undo").onclick = () => { zones[active].pop(); redraw(); };
$("clear").onclick = () => { zones[active] = []; redraw(); };

cv.addEventListener("click", e => {
  const r = cv.getBoundingClientRect();
  const x = (e.clientX - r.left) * cv.width / r.width;
  const y = (e.clientY - r.top) * cv.height / r.height;
  zones[active].push([Math.round(x), Math.round(y)]);
  redraw();
});

function redraw() {
  ctx.drawImage(img, 0, 0);
  const s = Math.max(1, cv.width / 900);
  for (const name of ["ENTRY_EXIT", "WORK_AREA"]) {
    const pts = zones[name]; if (!pts.length) continue;
    ctx.beginPath(); pts.forEach((p, i) => i ? ctx.lineTo(p[0], p[1]) : ctx.moveTo(p[0], p[1]));
    if (pts.length > 2) { ctx.closePath(); ctx.fillStyle = COLORS[name] + "44"; ctx.fill(); }
    ctx.strokeStyle = COLORS[name]; ctx.lineWidth = 3 * s; ctx.stroke();
    pts.forEach(p => { ctx.beginPath(); ctx.arc(p[0], p[1], 5 * s, 0, 7);
                       ctx.fillStyle = COLORS[name]; ctx.fill(); });
    ctx.font = `bold ${16 * s}px sans-serif`; ctx.fillStyle = COLORS[name];
    ctx.fillText(LABELS[name], pts[0][0] + 8 * s, pts[0][1] - 8 * s);
  }
  const ready = zones.ENTRY_EXIT.length >= 3 && zones.WORK_AREA.length >= 3;
  $("go").disabled = !ready;
  $("goHint").textContent = ready ? "Ready." : "Draw both areas to continue (3+ points each).";
}

/* ---------- Step 3: start + poll ---------- */
$("go").onclick = async () => {
  $("goErr").textContent = "";
  const body = {
    zones, work_threshold: (+$("thr").value) / 100,
    min_seconds: +$("minsec").value, tracker: $("trk").value,
    snapshots: $("snap").checked
  };
  const r = await fetch(`/start/${jobId}`, { method: "POST",
      headers: { "Content-Type": "application/json" }, body: JSON.stringify(body) });
  const j = await r.json();
  if (!r.ok) { $("goErr").textContent = j.error || "Could not start"; return; }
  $("s1").classList.add("hidden"); $("s2").classList.add("hidden");
  $("s3").classList.remove("hidden"); $("s3").scrollIntoView({behavior:"smooth"});
  timer = setInterval(poll, 800);
};

async function poll() {
  const j = await (await fetch(`/status/${jobId}`)).json();
  $("pb").style.width = (j.progress || 0) + "%";
  $("msg").textContent = (j.message || "") + (j.progress ? ` (${j.progress}%)` : "");
  $("live").textContent = (j.people_seen !== null && j.state === "running")
      ? `In frame now: ${j.people_now || 0}  |  Different people tracked so far: ${j.people_seen || 0}` : "";
  if (j.preview_ts !== null && j.preview_ts !== undefined && j.state === "running") {
    $("prev").classList.remove("hidden");
    $("prev").src = `/jobs/${jobId}/preview.jpg?t=${Date.now()}`;
  }
  if (j.state === "done") { clearInterval(timer); showResults(j); }
  if (j.state === "error") { clearInterval(timer); $("msg").textContent = "Error: " + j.message; }
}

/* ---------- Results ---------- */
function timeline(p, dur) {
  const d = Math.max(dur, 0.001);
  return '<div class="tl">' + p.segments.map(s => {
    const left = s.start / d * 100, width = Math.max((s.end - s.start) / d * 100, 0.4);
    return `<span style="left:${left.toFixed(2)}%;width:${width.toFixed(2)}%;background:${ZH[s.zone]}" title="${esc(s.zone_label)} ${s.start_fmt}-${s.end_fmt}"></span>`;
  }).join("") + '</div>';
}

function personCard(p, v) {
  const cls = p.verdict === "WORKING" ? "w" : p.verdict === "NOT WORKING" ? "n" : "b";
  const photo = p.photo ? `<img src="/jobs/${jobId}/${p.photo}?t=${Date.now()}" alt="${esc(p.name)}">`
                        : '<div class="noph">No photo</div>';
  const note = p.brief ? `<div class="hint" style="margin:0 0 8px">Seen for only ${p.total_s.toFixed(1)} s &ndash; too short to judge (may be a false detection or a split ID).</div>` : "";
  const kv = [
    ["Seen", `${p.first_seen} &ndash; ${p.last_seen}`], ["On screen", p.total_time],
    ["In work area", `${p.work_time} (${p.work_percent}%)`],
    ["In entry/exit area", p.entry_exit_time], ["Elsewhere", p.other_time],
    ["Work sessions", p.work_sessions], ["Longest work stretch", p.longest_work],
    ["Zone changes", p.zone_changes], ["How they appeared", esc(p.entered_via)],
    ["How they ended", esc(p.ended)], ["Zone path", esc(p.zone_path)],
    ["Movement", `${p.distance_px} px`],
    ["Detection confidence", `${p.avg_confidence} (${p.detections} frames)`]
  ].map(([k, val]) => `<div><span>${k}</span>${val}</div>`).join("");
  const rows = p.segments.map(s => `<tr><td>${esc(s.zone_label)}</td><td>${s.start_fmt}</td><td>${s.end_fmt}</td><td>${s.duration_fmt}</td></tr>`).join("");
  return `<div class="card ${p.brief ? "brief" : ""}">
    <div class="ph">${photo}</div>
    <div class="cbody">
      <div class="chead"><h4>${esc(p.name)}</h4><span class="badge ${cls}">${p.verdict}</span></div>
      ${note}<div class="kv">${kv}</div>
      ${timeline(p, v.duration_s)}
      <div class="scale"><span>00:00</span><span>${fmtTs(v.duration_s)}</span></div>
      <details><summary>Zone timeline (${p.segments.length} stays)</summary>
        <table class="seg"><tr><th>Area</th><th>From</th><th>To</th><th>Duration</th></tr>${rows}</table></details>
    </div></div>`;
}

function fmtTs(sec) {
  sec = Math.floor(sec); const h = Math.floor(sec / 3600), m = Math.floor(sec % 3600 / 60), s = sec % 60;
  const p2 = n => String(n).padStart(2, "0");
  return h ? `${h}:${p2(m)}:${p2(s)}` : `${p2(m)}:${p2(s)}`;
}

function renderPeople() {
  const f = $("flt").value;
  const list = RES.people.filter(p => f === "all" || p.verdict === f);
  $("people").innerHTML = list.map(p => personCard(p, RES.video)).join("");
  $("noRows").style.display = RES.people.length ? "none" : "block";
}
$("flt").onchange = () => { if (RES) renderPeople(); };

function showResults(j) {
  RES = j.result;
  $("s3").classList.add("hidden"); $("s4").classList.remove("hidden");
  const v = RES.video, s = RES.summary, st = RES.settings;
  const info = [
    ["Video", esc(v.name)], ["Duration", v.duration],
    ["Resolution", `${v.width} &times; ${v.height} @ ${v.fps} fps`], ["Analysed", v.analysed_at],
    ["Processing", `${v.processing_s}s on ${esc(v.device)}`], ["Tracker", esc(st.tracker)],
    ["Working rule", `&ge; ${st.work_threshold_percent}% of time in work area`],
    ["Brief if seen less than", `${st.brief_seconds} s`]
  ];
  $("vinfo").innerHTML = info.map(([k, val]) => `<div><span>${k}</span>${val}</div>`).join("");
  $("tiles").innerHTML = [
    [s.tracked_ids, "IDs tracked"], [s.working, "Working"], [s.not_working, "Not working"],
    [s.brief, "Brief detections"], [s.avg_work_percent + "%", "Avg work time (judged)"]
  ].map(([n, l]) => `<div class="stat"><b>${n}</b><span>${l}</span></div>`).join("");
  renderPeople();

  $("dl").innerHTML = (j.downloads || []).map(d => {
    const href = `/jobs/${jobId}/${d.file}` + (d.newtab ? "" : "?dl=1");
    return `<a class="btn" href="${href}" ${d.newtab ? 'target="_blank" rel="noopener"' : ""}>${esc(d.label)}<small>${esc(d.note || "")}</small></a>`;
  }).join("");

  if (j.video_playable) { $("vid").src = `/jobs/${jobId}/processed.mp4`; $("vid").classList.remove("hidden"); }
  else { $("noPlay").classList.remove("hidden"); }
  $("s4").scrollIntoView({behavior:"smooth"});
}
$("again").onclick = () => location.reload();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)